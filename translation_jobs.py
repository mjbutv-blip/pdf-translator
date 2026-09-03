import io
import json
import os
import tempfile
import threading
import time
import traceback
import zipfile
from datetime import datetime, timedelta
from uuid import uuid4

import openpyxl
from openai import OpenAI

from ai_client import ai_call_context, error_metadata
from config import (
    DEFAULT_FONT,
    OPENAI_API_KEY,
    PDF_JOB_HEARTBEAT_SECONDS,
    PDF_JOB_MAX_ATTEMPTS,
    PDF_JOB_STALE_SECONDS,
    WORKER_STALE_SECONDS,
)
from db import get_db_connection
from translation_core import (
    _now_iso,
    add_translated_textboxes_to_excel,
    build_scope_report_xlsx,
    get_customer_glossary_bytes_for_translation,
    init_db,
    load_glossary,
    run_pdf_translation,
    run_excel_translation,
)

PDF_JOB_CANCELLED_ERROR = "USER_CANCELLED"
WORKER_LOST_ERROR = "WORKER_LOST"
MAX_ATTEMPTS_EXCEEDED_ERROR = "MAX_ATTEMPTS_EXCEEDED"
OWNERSHIP_LOST_ERROR = "WORKER_OWNERSHIP_LOST"

# Development compatibility only. External worker mode does not use this map.
_RUNNING_JOB_THREADS: dict[str, threading.Thread] = {}


class WorkerOwnershipLost(RuntimeError):
    pass


def _log_worker_event(worker_id: str, event: str, job_id: str = "", status: str = "", **extra) -> None:
    payload = {
        "worker_id": worker_id,
        "job_id": job_id,
        "event": event,
        "status": status,
        **{k: v for k, v in extra.items() if v is not None},
    }
    print(json.dumps(payload, ensure_ascii=False, sort_keys=True))


def _parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def _result_meta_with_error(error: str, message: str = "") -> str:
    return json.dumps(
        {
            "error": error,
            "worker_error": error,
            "message": message,
            "candidate_ids_reliable": False,
        },
        ensure_ascii=False,
    )


def _truncate(value: str, limit: int = 2000) -> str:
    text = str(value or "")
    return text if len(text) <= limit else text[:limit] + "...[truncated]"


def _structured_failure_meta(exc: Exception, *, step: str, job_type: str) -> str:
    meta = error_metadata(exc, step=step)
    meta.update({
        "error": str(exc),
        "job_type": job_type,
        "traceback": _truncate(traceback.format_exc(), 4000),
        "candidate_ids": [],
        "candidate_count": 0,
        "candidate_ids_reliable": False,
    })
    return json.dumps(meta, ensure_ascii=False)


def _user_error_message(exc: Exception, *, default: str = "翻译失败") -> str:
    meta = error_metadata(exc)
    labels = {
        "AI_TIMEOUT": "AI 请求超时",
        "AI_RATE_LIMIT": "AI 服务限流，请稍后重试",
        "AI_NETWORK": "AI 网络连接失败",
        "AI_SERVER_ERROR": "AI 服务暂时不可用",
        "AI_INVALID_RESPONSE": "AI 返回内容无法解析",
        "AI_AUTH": "AI 认证失败，请检查 API Key",
        "AI_MODEL_UNAVAILABLE": "AI 模型不可用",
        "AI_INVALID_REQUEST": "AI 请求无效",
        "AI_CANCELLED": "任务已取消",
    }
    label = labels.get(meta.get("error_code"))
    return f"{default}：{label}" if label else default


def _set_job_step(job_id: str, worker_id: str, execution_mode: str, step: str, message: str) -> None:
    update_translation_job_owned(
        job_id,
        worker_id,
        execution_mode=execution_mode,
        message=message,
        result_meta=json.dumps({"current_step": step}, ensure_ascii=False),
    )


def _row_to_dict(row) -> dict | None:
    return dict(row) if row is not None else None


def create_translation_job(
    job_type: str,
    username: str,
    customer_id: str,
    source_file_name: str,
    input_bytes: bytes,
    aux_bytes: bytes | None,
    config: dict,
    *,
    execution_mode: str = "external",
    initial_status: str = "queued",
    initial_progress: float = 0.0,
    initial_message: str | None = None,
) -> str:
    if execution_mode not in {"external", "sync", "embedded"}:
        raise ValueError(f"Unsupported execution_mode: {execution_mode}")
    if initial_status not in {"queued", "running", "complete", "failed"}:
        raise ValueError(f"Unsupported initial_status: {initial_status}")
    job_id = str(uuid4())
    now = _now_iso()
    message = initial_message if initial_message is not None else ("等待开始" if initial_status == "queued" else "")
    with get_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO translation_jobs (
                job_id, job_type, status, username, customer_id, source_file_name,
                progress, message, input_bytes, aux_bytes, config,
                execution_mode, worker_id, heartbeat_at, attempt_count, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, NULL, 0, ?, ?)
            """,
            (
                job_id,
                job_type,
                initial_status,
                username,
                customer_id,
                source_file_name,
                initial_progress,
                message,
                input_bytes,
                aux_bytes,
                json.dumps(config, ensure_ascii=False),
                execution_mode,
                now,
                now,
            ),
        )
    return job_id


def create_external_excel_translation_job(
    *,
    username: str,
    customer_id: str,
    source_file_name: str,
    input_bytes: bytes,
    config: dict,
) -> str:
    return create_translation_job(
        job_type="Excel",
        username=username,
        customer_id=customer_id,
        source_file_name=source_file_name,
        input_bytes=input_bytes,
        aux_bytes=None,
        config=config,
        execution_mode="external",
    )


def update_translation_job(job_id: str, **fields) -> None:
    if not fields:
        return
    fields["updated_at"] = _now_iso()
    assignments = ", ".join(f"{key} = ?" for key in fields)
    values = list(fields.values()) + [job_id]
    with get_db_connection() as conn:
        conn.execute(
            f"UPDATE translation_jobs SET {assignments} WHERE job_id = ?",
            values,
        )


def update_translation_job_owned(
    job_id: str,
    worker_id: str,
    *,
    execution_mode: str | None = "external",
    **fields,
) -> bool:
    if not fields:
        return True
    fields["updated_at"] = _now_iso()
    assignments = ", ".join(f"{key} = ?" for key in fields)
    values = list(fields.values()) + [job_id, worker_id]
    mode_clause = ""
    if execution_mode is not None:
        mode_clause = " AND execution_mode = ?"
        values.append(execution_mode)
    with get_db_connection() as conn:
        cur = conn.execute(
            f"""
            UPDATE translation_jobs
            SET {assignments}
            WHERE job_id = ? AND worker_id = ? AND status = 'running'{mode_clause}
            """,
            values,
        )
        return bool(getattr(cur, "rowcount", 0) == 1)


def is_translation_job_cancelled(job_id: str) -> bool:
    with get_db_connection() as conn:
        row = conn.execute(
            "SELECT status, error FROM translation_jobs WHERE job_id = ?",
            (job_id,),
        ).fetchone()
    return bool(row and row["status"] == "failed" and row["error"] == PDF_JOB_CANCELLED_ERROR)


def raise_if_translation_job_cancelled(job_id: str) -> None:
    if is_translation_job_cancelled(job_id):
        raise RuntimeError(PDF_JOB_CANCELLED_ERROR)


def cancel_translation_jobs(username: str, job_type: str) -> dict:
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT job_id, status
            FROM translation_jobs
            WHERE username = ? AND job_type = ? AND status IN ('queued', 'running')
            """,
            (username, job_type),
        ).fetchall()
        job_ids = [row["job_id"] for row in rows]
        if job_ids:
            placeholders = ",".join("?" for _ in job_ids)
            conn.execute(
                f"""
                UPDATE translation_jobs
                SET status = 'failed',
                    progress = 0,
                    message = '用户已取消，可重新上传文件开始翻译',
                    error = ?,
                    worker_id = NULL,
                    heartbeat_at = NULL,
                    updated_at = ?
                WHERE job_id IN ({placeholders})
                """,
                [PDF_JOB_CANCELLED_ERROR, _now_iso(), *job_ids],
            )
    return {
        "cancelled_count": len(job_ids),
        "queued_count": sum(1 for row in rows if row["status"] == "queued"),
        "running_count": sum(1 for row in rows if row["status"] == "running"),
        "running_threads": [],
    }


def cancel_pdf_translation_jobs(username: str) -> dict:
    return cancel_translation_jobs(username, "PDF")


def cancel_excel_translation_jobs(username: str) -> dict:
    return cancel_translation_jobs(username, "Excel")


def list_translation_jobs(username: str, job_type: str = "PDF", limit: int = 20) -> list[dict]:
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT job_id, source_file_name, job_type, status, progress, message, error,
                   created_at, updated_at, execution_mode, worker_id, heartbeat_at
            FROM translation_jobs
            WHERE username = ? AND job_type = ?
            ORDER BY CASE status WHEN 'running' THEN 0 WHEN 'queued' THEN 1 ELSE 2 END,
                     created_at DESC
            LIMIT ?
            """,
            (username, job_type, limit),
        ).fetchall()
    return [dict(r) for r in rows]


def get_translation_job(job_id: str, username: str) -> dict | None:
    with get_db_connection() as conn:
        row = conn.execute(
            """
            SELECT *
            FROM translation_jobs
            WHERE job_id = ? AND username = ?
            """,
            (job_id, username),
        ).fetchone()
    return _row_to_dict(row)


def get_translation_job_result(job_id: str, username: str) -> dict | None:
    """Load result bytes for exactly one authorized completed job."""
    with get_db_connection() as conn:
        row = conn.execute(
            """
            SELECT job_id, source_file_name, job_type, status,
                   result_file, result_report, result_meta
            FROM translation_jobs
            WHERE job_id = ? AND username = ? AND status = 'complete'
            """,
            (job_id, username),
        ).fetchone()
    return _row_to_dict(row)


def get_translation_job_by_id(job_id: str) -> dict | None:
    with get_db_connection() as conn:
        row = conn.execute("SELECT * FROM translation_jobs WHERE job_id = ?", (job_id,)).fetchone()
    return _row_to_dict(row)


def register_translation_worker(worker_id: str) -> None:
    now = _now_iso()
    with get_db_connection() as conn:
        conn.execute(
            """
            INSERT INTO translation_workers (
                worker_id, started_at, heartbeat_at, status, stopped_at
            )
            VALUES (?, ?, ?, 'running', NULL)
            ON CONFLICT(worker_id) DO UPDATE SET
                started_at = excluded.started_at,
                heartbeat_at = excluded.heartbeat_at,
                status = 'running',
                stopped_at = NULL
            """,
            (worker_id, now, now),
        )


def heartbeat_translation_worker(worker_id: str) -> bool:
    now = _now_iso()
    with get_db_connection() as conn:
        cur = conn.execute(
            """
            UPDATE translation_workers
            SET heartbeat_at = ?,
                status = 'running',
                stopped_at = NULL
            WHERE worker_id = ?
            """,
            (now, worker_id),
        )
    return bool(getattr(cur, "rowcount", 0) == 1)


def stop_translation_worker(worker_id: str) -> None:
    now = _now_iso()
    with get_db_connection() as conn:
        conn.execute(
            """
            UPDATE translation_workers
            SET heartbeat_at = ?,
                status = 'stopped',
                stopped_at = ?
            WHERE worker_id = ?
            """,
            (now, now, worker_id),
        )


def list_live_workers(*, now: datetime | None = None) -> list[dict]:
    health = get_worker_health(now=now, include_workers=True)
    return [row for row in health.get("workers", []) if row.get("is_live")]


def get_worker_health(*, now: datetime | None = None, include_workers: bool = False) -> dict:
    now_dt = now or datetime.now()
    stale_before = now_dt - timedelta(seconds=WORKER_STALE_SECONDS)
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT worker_id, started_at, heartbeat_at, status, stopped_at
            FROM translation_workers
            ORDER BY heartbeat_at DESC, started_at DESC
            """
        ).fetchall()
    live_count = 0
    stale_count = 0
    latest_heartbeat_at = None
    worker_rows: list[dict] = []
    for row in rows:
        item = dict(row)
        heartbeat_dt = _parse_iso(item.get("heartbeat_at"))
        is_running = item.get("status") == "running"
        is_live = bool(is_running and heartbeat_dt and heartbeat_dt >= stale_before)
        is_stale = bool(is_running and (not heartbeat_dt or heartbeat_dt < stale_before))
        if is_live:
            live_count += 1
        if is_stale:
            stale_count += 1
        if heartbeat_dt and (latest_heartbeat_at is None or heartbeat_dt > latest_heartbeat_at):
            latest_heartbeat_at = heartbeat_dt
        item["is_live"] = is_live
        item["is_stale"] = is_stale
        worker_rows.append(item)
    result = {
        "live_worker_count": live_count,
        "stale_worker_count": stale_count,
        "latest_heartbeat_at": latest_heartbeat_at.isoformat() if latest_heartbeat_at else None,
        "worker_stale_seconds": WORKER_STALE_SECONDS,
    }
    if include_workers:
        result["workers"] = worker_rows
    return result


def get_translation_queue_health(*, now: datetime | None = None) -> dict:
    now_dt = now or datetime.now()
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT job_type, status, MIN(created_at) AS oldest_created_at, COUNT(*) AS n
            FROM translation_jobs
            WHERE execution_mode = 'external'
              AND status IN ('queued', 'running')
              AND job_type IN ('PDF', 'Excel')
            GROUP BY job_type, status
            """
        ).fetchall()
    summary = {
        "queued_count": 0,
        "running_count": 0,
        "queued_pdf_count": 0,
        "queued_excel_count": 0,
        "running_pdf_count": 0,
        "running_excel_count": 0,
        "oldest_queued_created_at": None,
        "oldest_queue_age_seconds": None,
    }
    oldest_dt = None
    for row in rows:
        job_type = row["job_type"]
        status = row["status"]
        count = int(row["n"] or 0)
        if status == "queued":
            summary["queued_count"] += count
            key = "queued_pdf_count" if job_type == "PDF" else "queued_excel_count"
            summary[key] += count
            candidate_oldest = _parse_iso(row["oldest_created_at"])
            if candidate_oldest and (oldest_dt is None or candidate_oldest < oldest_dt):
                oldest_dt = candidate_oldest
        elif status == "running":
            summary["running_count"] += count
            key = "running_pdf_count" if job_type == "PDF" else "running_excel_count"
            summary[key] += count
    if oldest_dt:
        summary["oldest_queued_created_at"] = oldest_dt.isoformat()
        summary["oldest_queue_age_seconds"] = max(0, int((now_dt - oldest_dt).total_seconds()))
    return summary


def get_worker_queue_health(*, now: datetime | None = None) -> dict:
    health = get_worker_health(now=now)
    queue = get_translation_queue_health(now=now)
    return {
        **health,
        **queue,
        "worker_available": health["live_worker_count"] > 0,
        "queue_waiting_without_worker": queue["queued_count"] > 0 and health["live_worker_count"] == 0,
    }


def delete_translation_job(job_id: str, username: str) -> None:
    with get_db_connection() as conn:
        conn.execute(
            """
            DELETE FROM translation_candidate_occurrences
            WHERE translation_job_id IN (
                SELECT job_id FROM translation_jobs WHERE job_id = ? AND username = ?
            )
            """,
            (job_id, username),
        )
        conn.execute(
            "DELETE FROM translation_jobs WHERE job_id = ? AND username = ?",
            (job_id, username),
        )


def list_term_candidates_for_job(translation_job_id: str) -> list[dict]:
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT
                tc.candidate_id,
                tc.customer_id,
                tc.original_term,
                tc.normalized_term,
                tc.ai_suggested_translation,
                tc.final_translation,
                tc.status,
                tc.confidence,
                tco.occurrence_count,
                tco.representative_page_or_sheet,
                tco.representative_cell_coordinate,
                tco.original_term_snapshot,
                tco.normalized_term_snapshot,
                tco.context_sentence_snapshot,
                tco.created_at AS occurrence_created_at
            FROM translation_candidate_occurrences tco
            JOIN term_candidates tc ON tc.candidate_id = tco.candidate_id
            WHERE tco.translation_job_id = ?
            ORDER BY tco.created_at ASC, tc.candidate_id ASC
            """,
            (translation_job_id,),
        ).fetchall()
    return [dict(r) for r in rows]


def list_candidate_occurrences(candidate_id: int) -> list[dict]:
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT
                tco.occurrence_id,
                tco.translation_job_id,
                tco.candidate_id,
                tj.source_file_name,
                tj.job_type,
                tj.customer_id,
                tj.created_at AS job_created_at,
                tj.status AS job_status,
                tco.occurrence_count,
                tco.representative_page_or_sheet,
                tco.representative_cell_coordinate,
                tco.original_term_snapshot,
                tco.normalized_term_snapshot,
                tco.context_sentence_snapshot,
                tco.created_at AS occurrence_created_at
            FROM translation_candidate_occurrences tco
            JOIN translation_jobs tj ON tj.job_id = tco.translation_job_id
            WHERE tco.candidate_id = ?
            ORDER BY tco.created_at DESC, tco.occurrence_id DESC
            """,
            (candidate_id,),
        ).fetchall()
    return [dict(r) for r in rows]


def get_next_queued_pdf_job(username: str, exclude_job_id: str = "") -> dict | None:
    with get_db_connection() as conn:
        row = conn.execute(
            """
            SELECT job_id
            FROM translation_jobs
            WHERE username = ?
              AND job_type = 'PDF'
              AND status = 'queued'
              AND execution_mode = 'embedded'
              AND job_id <> ?
            ORDER BY created_at ASC
            LIMIT 1
            """,
            (username, exclude_job_id),
        ).fetchone()
    return _row_to_dict(row)


def _supported_external_job_types() -> tuple[str, ...]:
    return ("PDF", "Excel")


def recover_stale_external_jobs(worker_id: str, *, now: datetime | None = None) -> dict:
    now_dt = now or datetime.now()
    stale_before = now_dt - timedelta(seconds=PDF_JOB_STALE_SECONDS)
    recovered: list[str] = []
    failed: list[str] = []
    with get_db_connection() as conn:
        rows = conn.execute(
            """
            SELECT job_id, job_type, worker_id, heartbeat_at, updated_at, attempt_count
            FROM translation_jobs
            WHERE status = 'running'
              AND execution_mode = 'external'
              AND worker_id IS NOT NULL
              AND job_type IN ('PDF', 'Excel')
            """
        ).fetchall()
        for row in rows:
            last_seen = _parse_iso(row["heartbeat_at"]) or _parse_iso(row["updated_at"])
            if last_seen and last_seen > stale_before:
                continue
            job_id = row["job_id"]
            old_worker_id = row["worker_id"]
            attempt_count = int(row["attempt_count"] or 0)
            if attempt_count >= PDF_JOB_MAX_ATTEMPTS:
                cur = conn.execute(
                    """
                    UPDATE translation_jobs
                    SET status = 'failed',
                        progress = 0,
                        message = '后台翻译服务多次中断，已停止重试',
                        error = ?,
                        result_meta = ?,
                        worker_id = NULL,
                        heartbeat_at = NULL,
                        updated_at = ?
                    WHERE job_id = ?
                      AND status = 'running'
                      AND execution_mode = 'external'
                      AND COALESCE(worker_id, '') = COALESCE(?, '')
                    """,
                    (
                        MAX_ATTEMPTS_EXCEEDED_ERROR,
                        _result_meta_with_error(MAX_ATTEMPTS_EXCEEDED_ERROR, "stale worker exceeded max attempts"),
                        _now_iso(),
                        job_id,
                        old_worker_id,
                    ),
                )
                if getattr(cur, "rowcount", 0) == 1:
                    failed.append(job_id)
                    _log_worker_event(worker_id, "job_recovery_failed_max_attempts", job_id, "failed", job_type=row["job_type"])
                continue
            cur = conn.execute(
                """
                UPDATE translation_jobs
                SET status = 'queued',
                    progress = 0,
                    message = '后台翻译服务中断，已重新排队',
                    error = ?,
                    worker_id = NULL,
                    heartbeat_at = NULL,
                    updated_at = ?
                WHERE job_id = ?
                  AND status = 'running'
                  AND execution_mode = 'external'
                  AND COALESCE(worker_id, '') = COALESCE(?, '')
                """,
                (
                    WORKER_LOST_ERROR,
                    _now_iso(),
                    job_id,
                    old_worker_id,
                ),
            )
            if getattr(cur, "rowcount", 0) == 1:
                recovered.append(job_id)
                _log_worker_event(worker_id, "job_recovered", job_id, "queued", job_type=row["job_type"])
    return {"recovered": recovered, "failed": failed}


def recover_stale_pdf_jobs(worker_id: str, *, now: datetime | None = None) -> dict:
    return recover_stale_external_jobs(worker_id, now=now)


def _claim_next_external_translation_job_sqlite(worker_id: str) -> dict | None:
    now = _now_iso()
    with get_db_connection() as conn:
        conn.conn.execute("BEGIN IMMEDIATE")
        row = conn.execute(
            """
            SELECT *
            FROM translation_jobs
            WHERE status = 'queued'
              AND execution_mode = 'external'
              AND job_type IN ('PDF', 'Excel')
            ORDER BY created_at ASC
            LIMIT 1
            """
        ).fetchone()
        if row is None:
            return None
        cur = conn.execute(
            """
            UPDATE translation_jobs
            SET status = 'running',
                worker_id = ?,
                heartbeat_at = ?,
                attempt_count = COALESCE(attempt_count, 0) + 1,
                progress = 0.01,
                message = '后台翻译服务已领取任务',
                error = '',
                updated_at = ?
            WHERE job_id = ?
              AND status = 'queued'
              AND execution_mode = 'external'
              AND job_type IN ('PDF', 'Excel')
            """,
            (worker_id, now, now, row["job_id"]),
        )
        if getattr(cur, "rowcount", 0) != 1:
            return None
        claimed = conn.execute("SELECT * FROM translation_jobs WHERE job_id = ?", (row["job_id"],)).fetchone()
    return _row_to_dict(claimed)


def _claim_next_external_translation_job_postgres(worker_id: str) -> dict | None:
    now = _now_iso()
    with get_db_connection() as conn:
        row = conn.execute(
            """
            UPDATE translation_jobs
            SET status = 'running',
                worker_id = ?,
                heartbeat_at = ?,
                attempt_count = COALESCE(attempt_count, 0) + 1,
                progress = 0.01,
                message = '后台翻译服务已领取任务',
                error = '',
                updated_at = ?
            WHERE job_id = (
                SELECT job_id
                FROM translation_jobs
                WHERE status = 'queued'
                  AND execution_mode = 'external'
                  AND job_type IN ('PDF', 'Excel')
                ORDER BY created_at ASC
                LIMIT 1
                FOR UPDATE SKIP LOCKED
            )
            RETURNING *
            """,
            (worker_id, now, now),
        ).fetchone()
    return _row_to_dict(row)


def claim_next_external_translation_job(worker_id: str) -> dict | None:
    init_db()
    recover_stale_external_jobs(worker_id)
    with get_db_connection() as conn:
        is_postgres = conn.is_postgres
    job = (
        _claim_next_external_translation_job_postgres(worker_id)
        if is_postgres
        else _claim_next_external_translation_job_sqlite(worker_id)
    )
    if job:
        _log_worker_event(
            worker_id,
            "job_claimed",
            job["job_id"],
            job["status"],
            job_type=job.get("job_type"),
            attempt_count=job.get("attempt_count"),
        )
    return job


def claim_next_pdf_job(worker_id: str) -> dict | None:
    return claim_next_external_translation_job(worker_id)


def heartbeat_pdf_job(job_id: str, worker_id: str, *, execution_mode: str = "external") -> bool:
    return update_translation_job_owned(
        job_id,
        worker_id,
        execution_mode=execution_mode,
        heartbeat_at=_now_iso(),
    )


def _start_worker_heartbeat(job_id: str, worker_id: str, execution_mode: str) -> tuple[threading.Event, threading.Thread]:
    stop_event = threading.Event()

    def heartbeat_loop() -> None:
        while not stop_event.wait(PDF_JOB_HEARTBEAT_SECONDS):
            if not heartbeat_pdf_job(job_id, worker_id, execution_mode=execution_mode):
                _log_worker_event(worker_id, "heartbeat_stopped", job_id)
                stop_event.set()
                return
            _log_worker_event(worker_id, "heartbeat", job_id, "running")

    thread = threading.Thread(
        target=heartbeat_loop,
        name=f"pdf-job-heartbeat-{job_id}",
        daemon=True,
    )
    thread.start()
    return stop_event, thread


def _assert_worker_still_owns(job_id: str, worker_id: str, execution_mode: str | None = None) -> None:
    row = get_translation_job_by_id(job_id)
    if not row:
        raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)
    if row.get("status") == "failed" and row.get("error") == PDF_JOB_CANCELLED_ERROR:
        raise RuntimeError(PDF_JOB_CANCELLED_ERROR)
    if row.get("status") != "running" or row.get("worker_id") != worker_id:
        raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)
    if execution_mode is not None and row.get("execution_mode") != execution_mode:
        raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)


def _pdf_scope_report_rows(file_name: str, scope_cfg: dict, selected_pages, scope_detection, scope_mode: str) -> list[dict]:
    selected_page_nums = [
        pn + 1 for pn in (selected_pages or [])
    ] if selected_pages is not None else list(range(1, int(scope_cfg.get("total_pages", 0)) + 1))
    detection_by_page = {
        row.get("page_number"): row for row in (scope_detection or [])
    }
    total_pages = int(scope_cfg.get("total_pages") or len(detection_by_page) or 0)
    rows = []
    if total_pages:
        selected_set = set(selected_page_nums)
        for page_number in range(1, total_pages + 1):
            det = detection_by_page.get(page_number, {})
            rows.append({
                "file_name": file_name,
                "source_type": "PDF",
                "scope_mode": scope_mode,
                "item": f"第 {page_number} 页",
                "selected": page_number in selected_set,
                "score": det.get("score", ""),
                "reason": det.get("reason", ""),
            })
    return rows


def _build_excel_report_bytes(report_rows: list[dict], img_report_rows: list[dict] | None = None) -> bytes:
    report_wb = openpyxl.Workbook()
    report_ws = report_wb.active
    report_ws.title = "翻译报告"
    headers = [
        "sheet_name", "cell_coordinate", "original_text", "translated_text",
        "status", "skip_reason", "is_merged_cell", "layout_warning",
        "scope_mode", "selected_sheets", "skipped_sheets",
        "detection_score", "detection_reason",
    ]
    report_ws.append(headers)
    for cell in report_ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for row in report_rows:
        report_ws.append([row.get(header, "") for header in headers])
    for col_letter, width in zip(
        "ABCDEFGHIJKLM",
        [12, 14, 35, 35, 12, 24, 14, 14, 18, 35, 35, 12, 70],
    ):
        report_ws.column_dimensions[col_letter].width = width

    if img_report_rows:
        img_ws = report_wb.create_sheet("图片译文")
        img_headers = ["drawing", "image", "status", "original_text", "translated_text", "skip_reason"]
        img_ws.append(img_headers)
        for cell in img_ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for row in img_report_rows:
            img_ws.append([row.get(header, "") for header in img_headers])
        for col_letter, width in zip("ABCDEF", [25, 18, 10, 35, 35, 20]):
            img_ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    report_wb.save(buf)
    return buf.getvalue()


def run_claimed_pdf_job(
    job: dict,
    api_key: str,
    worker_id: str,
    *,
    translator=run_pdf_translation,
) -> str:
    job_id = job["job_id"]
    execution_mode = job.get("execution_mode") or "external"
    job_started = time.monotonic()
    step = {"name": "load_input"}
    _log_worker_event(
        worker_id, "job_started", job_id, "running", job_type="PDF", step=step["name"],
        filename=job.get("source_file_name"), progress=job.get("progress"),
        attempt_count=job.get("attempt_count"), heartbeat_at=job.get("heartbeat_at"),
    )
    heartbeat_stop, heartbeat_thread = _start_worker_heartbeat(job_id, worker_id, execution_mode)
    try:
        _assert_worker_still_owns(job_id, worker_id, execution_mode)
        config = json.loads(job.get("config") or "{}")
        last_progress_write = {"ts": 0.0, "value": 0.0}
        last_heartbeat = {"ts": 0.0}

        def heartbeat_if_due(force: bool = False) -> None:
            now = time.monotonic()
            if force or now - last_heartbeat["ts"] >= PDF_JOB_HEARTBEAT_SECONDS:
                last_heartbeat["ts"] = now
                if not heartbeat_pdf_job(job_id, worker_id, execution_mode=execution_mode):
                    raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

        font_bytes = job.get("aux_bytes")
        with tempfile.TemporaryDirectory() as tmpdir:
            step["name"] = "load_input"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在读取 PDF 输入")
            if font_bytes:
                font_path = os.path.join(tmpdir, "font.ttf")
                with open(font_path, "wb") as f:
                    f.write(font_bytes)
            else:
                font_path = str(DEFAULT_FONT)

            step["name"] = "load_glossary"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在加载客户术语库")
            glossary_bytes = get_customer_glossary_bytes_for_translation(
                {"username": job["username"], "role": "company_admin"},
                job["customer_id"],
            )

            def on_page(pn, total, n_blocks):
                _assert_worker_still_owns(job_id, worker_id, execution_mode)
                heartbeat_if_due()
                if not update_translation_job_owned(
                    job_id,
                    worker_id,
                    execution_mode=execution_mode,
                    heartbeat_at=_now_iso(),
                    message=f"第 {pn + 1}/{total} 页，{n_blocks} 个文本块",
                ):
                    raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

            def on_block(preview):
                _assert_worker_still_owns(job_id, worker_id, execution_mode)
                heartbeat_if_due()
                if not update_translation_job_owned(
                    job_id,
                    worker_id,
                    execution_mode=execution_mode,
                    heartbeat_at=_now_iso(),
                    message=str(preview)[:180],
                ):
                    raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

            def on_progress(frac):
                _assert_worker_still_owns(job_id, worker_id, execution_mode)
                progress = max(0.0, min(float(frac), 1.0))
                now = time.monotonic()
                should_write = (
                    progress >= 1.0
                    or progress - last_progress_write["value"] >= 0.01
                    or now - last_progress_write["ts"] >= 1.5
                )
                if should_write:
                    last_progress_write["ts"] = now
                    last_progress_write["value"] = progress
                    if not update_translation_job_owned(
                        job_id,
                        worker_id,
                        execution_mode=execution_mode,
                        heartbeat_at=_now_iso(),
                        progress=progress,
                    ):
                        raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)
                else:
                    heartbeat_if_due()

            selected_pages = config.get("selected_pages")
            scope_detection = config.get("scope_detection") or []
            scope_cfg = config.get("scope_cfg") or {}
            scope_mode = config.get("scope_mode") or "all"
            step["name"] = "translate_text"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在翻译 PDF 文本")
            with ai_call_context(
                job_id=job_id,
                job_type="PDF",
                step=step["name"],
                cancel_check=lambda: is_translation_job_cancelled(job_id),
            ):
                pdf_out, xlsx_out, n_terms, review_summary = translator(
                    pdf_bytes=job["input_bytes"],
                    glossary_bytes=glossary_bytes,
                    font_path=font_path,
                    api_key=api_key,
                    on_page=on_page,
                    on_block=on_block,
                    on_progress=on_progress,
                    customer_id=job["customer_id"],
                    source_file_name=job["source_file_name"],
                    created_by=job["username"],
                    selected_pages=selected_pages,
                    scope_mode=scope_mode,
                    scope_detection=scope_detection,
                    translation_job_id=job_id,
                )
            _assert_worker_still_owns(job_id, worker_id, execution_mode)
            step["name"] = "build_output"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在生成翻译结果")
            scope_report = _pdf_scope_report_rows(
                job["source_file_name"],
                scope_cfg,
                selected_pages,
                scope_detection,
                scope_mode,
            )
            scope_report_bytes = build_scope_report_xlsx(scope_report) if scope_report else b""
            report_bytes = b""
            report_kind = ""
            if xlsx_out and scope_report_bytes:
                report_zip = io.BytesIO()
                with zipfile.ZipFile(report_zip, "w") as zf:
                    zf.writestr("unrecorded_terms.xlsx", xlsx_out)
                    zf.writestr("scope_report.xlsx", scope_report_bytes)
                report_bytes = report_zip.getvalue()
                report_kind = "zip"
            elif xlsx_out:
                report_bytes = xlsx_out
                report_kind = "unrecorded"
            elif scope_report_bytes:
                report_bytes = scope_report_bytes
                report_kind = "scope"
            step["name"] = "finalize"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在保存翻译结果")
            job_candidate_rows = list_term_candidates_for_job(job_id)
            job_candidate_ids = [int(row["candidate_id"]) for row in job_candidate_rows]
            completed = update_translation_job_owned(
                job_id,
                worker_id,
                execution_mode=execution_mode,
                status="complete",
                progress=1.0,
                message="翻译完成",
                error="",
                result_file=pdf_out,
                result_report=report_bytes,
                result_meta=json.dumps(
                    {
                        "n_terms": n_terms,
                        "unrecorded_term_count": n_terms,
                        "unrecorded_terms": list((review_summary or {}).get("unrecorded_terms") or []),
                        "has_scope_report": bool(scope_report),
                        "has_unrecorded_terms": bool(xlsx_out),
                        "report_kind": report_kind,
                        "review_summary": review_summary,
                        "candidate_ids": job_candidate_ids,
                        "candidate_count": len(job_candidate_ids),
                        "candidate_ids_reliable": True,
                        "review_item_count": int((review_summary or {}).get("n_review_items", 0) or 0),
                        "duration_ms": int((time.monotonic() - job_started) * 1000),
                    },
                    ensure_ascii=False,
                ),
                heartbeat_at=_now_iso(),
            )
            if not completed:
                raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)
            _log_worker_event(
                worker_id,
                "job_completed",
                job_id,
                "complete",
                job_type="PDF",
                step=step["name"],
                duration_ms=int((time.monotonic() - job_started) * 1000),
            )
            return "complete"
    except Exception as exc:
        if str(exc) == PDF_JOB_CANCELLED_ERROR:
            _log_worker_event(worker_id, "job_cancelled", job_id, "failed")
            return "cancelled"
        if isinstance(exc, WorkerOwnershipLost):
            _log_worker_event(worker_id, "job_ownership_lost", job_id, "", error=str(exc))
            return "ownership_lost"
        failed = update_translation_job_owned(
            job_id,
            worker_id,
            execution_mode=execution_mode,
            status="failed",
            error=str(exc),
            message=_user_error_message(exc),
            result_meta=_structured_failure_meta(exc, step=step["name"], job_type="PDF"),
            heartbeat_at=_now_iso(),
        )
        err_meta = error_metadata(exc, step=step["name"])
        _log_worker_event(
            worker_id,
            "job_failed",
            job_id,
            "failed" if failed else "",
            job_type="PDF",
            step=step["name"],
            duration_ms=int((time.monotonic() - job_started) * 1000),
            error_code=err_meta.get("error_code"),
            retryable=err_meta.get("retryable"),
        )
        return "failed"
    finally:
        heartbeat_stop.set()
        heartbeat_thread.join(timeout=1)


def run_claimed_excel_job(
    job: dict,
    api_key: str,
    worker_id: str,
    *,
    translator=run_excel_translation,
    image_translator=add_translated_textboxes_to_excel,
) -> str:
    job_id = job["job_id"]
    execution_mode = job.get("execution_mode") or "external"
    job_started = time.monotonic()
    step = {"name": "load_input"}
    _log_worker_event(
        worker_id, "job_started", job_id, "running", job_type="Excel", step=step["name"],
        filename=job.get("source_file_name"), progress=job.get("progress"),
        attempt_count=job.get("attempt_count"), heartbeat_at=job.get("heartbeat_at"),
    )
    heartbeat_stop, heartbeat_thread = _start_worker_heartbeat(job_id, worker_id, execution_mode)
    try:
        _assert_worker_still_owns(job_id, worker_id, execution_mode)
        config = json.loads(job.get("config") or "{}")
        last_progress_write = {"ts": 0.0, "value": 0.0}

        step["name"] = "load_glossary"
        _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在加载客户术语库")
        glossary_bytes = get_customer_glossary_bytes_for_translation(
            {"username": job["username"], "role": "company_admin"},
            job["customer_id"],
        )

        def on_cell(preview):
            _assert_worker_still_owns(job_id, worker_id, execution_mode)
            if not update_translation_job_owned(
                job_id,
                worker_id,
                execution_mode=execution_mode,
                heartbeat_at=_now_iso(),
                message=str(preview)[:180],
            ):
                raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

        def on_progress(frac):
            _assert_worker_still_owns(job_id, worker_id, execution_mode)
            progress = max(0.0, min(float(frac), 1.0))
            now = time.monotonic()
            should_write = (
                progress >= 1.0
                or progress - last_progress_write["value"] >= 0.01
                or now - last_progress_write["ts"] >= 1.5
            )
            if should_write:
                last_progress_write["ts"] = now
                last_progress_write["value"] = progress
                if not update_translation_job_owned(
                    job_id,
                    worker_id,
                    execution_mode=execution_mode,
                    heartbeat_at=_now_iso(),
                    progress=progress,
                ):
                    raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

        selected_sheets = config.get("selected_sheets")
        scope_detection = config.get("scope_detection") or []
        scope_mode = config.get("scope_mode") or "all"
        translate_images = bool(config.get("translate_images"))
        step["name"] = "translate_text"
        _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在翻译 Excel 单元格")
        with ai_call_context(
            job_id=job_id,
            job_type="Excel",
            step=step["name"],
            cancel_check=lambda: is_translation_job_cancelled(job_id),
        ):
            excel_out, n_cells, _n_images, report_rows, review_summary = translator(
                xlsx_bytes=job["input_bytes"],
                glossary_bytes=glossary_bytes,
                api_key=api_key,
                on_cell=on_cell,
                on_progress=on_progress,
                translate_images=False,
                customer_id=job["customer_id"],
                source_file_name=job["source_file_name"],
                created_by=job["username"],
                selected_sheets=selected_sheets,
                scope_mode=scope_mode,
                scope_detection=scope_detection,
                translation_job_id=job_id,
            )
        _assert_worker_still_owns(job_id, worker_id, execution_mode)
        img_report_rows: list[dict] = []
        if translate_images:
            step["name"] = "translate_images"
            _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在翻译 Excel 图片文字")
            glossary_dict = load_glossary(glossary_bytes)
            client = OpenAI(api_key=api_key)

            def on_image(i, total, fname):
                _assert_worker_still_owns(job_id, worker_id, execution_mode)
                progress = 0.6 + (float(i) / max(float(total), 1.0)) * 0.4
                if not update_translation_job_owned(
                    job_id,
                    worker_id,
                    execution_mode=execution_mode,
                    heartbeat_at=_now_iso(),
                    message=f"图片 {i}/{total}：{fname}",
                    progress=progress,
                ):
                    raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)

            with ai_call_context(
                job_id=job_id,
                job_type="Excel",
                step=step["name"],
                cancel_check=lambda: is_translation_job_cancelled(job_id),
            ):
                excel_out, img_report_rows = image_translator(
                    excel_out,
                    client,
                    glossary_dict,
                    on_image,
                    selected_sheets=selected_sheets,
                )

        _assert_worker_still_owns(job_id, worker_id, execution_mode)
        step["name"] = "build_report"
        _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在生成 Excel 翻译报告")
        report_bytes = _build_excel_report_bytes(report_rows, img_report_rows)
        translated_image_count = sum(1 for row in img_report_rows if row.get("status") == "ok")
        job_candidate_rows = list_term_candidates_for_job(job_id)
        job_candidate_ids = [int(row["candidate_id"]) for row in job_candidate_rows]
        unrecorded_terms = list((review_summary or {}).get("unrecorded_terms") or [])
        unrecorded_term_count = int((review_summary or {}).get("unrecorded_term_count", len(unrecorded_terms)) or 0)
        step["name"] = "finalize"
        _set_job_step(job_id, worker_id, execution_mode, step["name"], "正在保存翻译结果")
        completed = update_translation_job_owned(
            job_id,
            worker_id,
            execution_mode=execution_mode,
            status="complete",
            progress=1.0,
            message="翻译完成",
            error="",
            result_file=excel_out,
            result_report=report_bytes,
            result_meta=json.dumps(
                {
                    "customer_id": job["customer_id"],
                    "selected_sheets": selected_sheets,
                    "scope_mode": scope_mode,
                    "translate_images": translate_images,
                    "translated_cell_count": n_cells,
                    "translated_image_count": translated_image_count,
                    "unrecorded_term_count": unrecorded_term_count,
                    "unrecorded_terms": unrecorded_terms,
                    "image_translation_unrecorded_tracking_supported": False,
                    "candidate_ids": job_candidate_ids,
                    "candidate_count": len(job_candidate_ids),
                    "candidate_ids_reliable": True,
                    "review_item_count": int((review_summary or {}).get("n_review_items", 0) or 0),
                    "review_summary": review_summary,
                    "duration_ms": int((time.monotonic() - job_started) * 1000),
                },
                ensure_ascii=False,
            ),
            heartbeat_at=_now_iso(),
        )
        if not completed:
            raise WorkerOwnershipLost(OWNERSHIP_LOST_ERROR)
        _log_worker_event(
            worker_id,
            "job_completed",
            job_id,
            "complete",
            job_type="Excel",
            step=step["name"],
            duration_ms=int((time.monotonic() - job_started) * 1000),
        )
        return "complete"
    except Exception as exc:
        if str(exc) == PDF_JOB_CANCELLED_ERROR:
            _log_worker_event(worker_id, "job_cancelled", job_id, "failed", job_type="Excel")
            return "cancelled"
        if isinstance(exc, WorkerOwnershipLost):
            _log_worker_event(worker_id, "job_ownership_lost", job_id, "", job_type="Excel", error=str(exc))
            return "ownership_lost"
        failed = update_translation_job_owned(
            job_id,
            worker_id,
            execution_mode=execution_mode,
            status="failed",
            error=str(exc),
            message=_user_error_message(exc),
            result_meta=_structured_failure_meta(exc, step=step["name"], job_type="Excel"),
            heartbeat_at=_now_iso(),
        )
        err_meta = error_metadata(exc, step=step["name"])
        _log_worker_event(
            worker_id,
            "job_failed",
            job_id,
            "failed" if failed else "",
            job_type="Excel",
            step=step["name"],
            duration_ms=int((time.monotonic() - job_started) * 1000),
            error_code=err_meta.get("error_code"),
            retryable=err_meta.get("retryable"),
        )
        return "failed"
    finally:
        heartbeat_stop.set()
        heartbeat_thread.join(timeout=1)


def run_claimed_translation_job(job: dict, api_key: str, worker_id: str) -> str:
    if job.get("job_type") == "PDF":
        return run_claimed_pdf_job(job, api_key, worker_id)
    if job.get("job_type") == "Excel":
        return run_claimed_excel_job(job, api_key, worker_id)
    update_translation_job_owned(
        job["job_id"],
        worker_id,
        execution_mode=job.get("execution_mode") or "external",
        status="failed",
        error=f"Unsupported job_type: {job.get('job_type')}",
        message="翻译失败",
    )
    return "failed"


def _run_pdf_translation_job(job_id: str, api_key: str, start_next_on_finish: bool = False) -> None:
    worker_id = f"embedded-{uuid4()}"
    job = claim_pdf_job_by_id(job_id, worker_id)
    if not job:
        return
    try:
        run_claimed_pdf_job(job, api_key, worker_id)
    finally:
        _RUNNING_JOB_THREADS.pop(job_id, None)
        if start_next_on_finish and not is_translation_job_cancelled(job_id):
            next_job_ref = get_next_queued_pdf_job(job["username"], exclude_job_id=job_id)
            if next_job_ref:
                next_job = claim_pdf_job_by_id(next_job_ref["job_id"], worker_id)
                if next_job:
                    run_claimed_pdf_job(next_job, api_key, worker_id)


def claim_pdf_job_by_id(job_id: str, worker_id: str) -> dict | None:
    now = _now_iso()
    with get_db_connection() as conn:
        cur = conn.execute(
            """
            UPDATE translation_jobs
            SET status = 'running',
                worker_id = ?,
                heartbeat_at = ?,
                attempt_count = COALESCE(attempt_count, 0) + 1,
                progress = 0.01,
                message = '后台翻译服务已领取任务',
                error = '',
                updated_at = ?
            WHERE job_id = ?
              AND job_type = 'PDF'
              AND status = 'queued'
              AND execution_mode = 'embedded'
            """,
            (worker_id, now, now, job_id),
        )
        if getattr(cur, "rowcount", 0) != 1:
            return None
        row = conn.execute("SELECT * FROM translation_jobs WHERE job_id = ?", (job_id,)).fetchone()
    return _row_to_dict(row)


def start_pdf_translation_job(job_id: str, api_key: str | None = None, start_next_on_finish: bool = False) -> bool:
    if os.getenv("PDF_WORKER_MODE", "external").strip().lower() != "embedded":
        return False
    if job_id in _RUNNING_JOB_THREADS and _RUNNING_JOB_THREADS[job_id].is_alive():
        return False
    resolved_api_key = api_key or OPENAI_API_KEY
    if not resolved_api_key:
        return False
    thread = threading.Thread(
        target=_run_pdf_translation_job,
        args=(job_id, resolved_api_key, start_next_on_finish),
        daemon=True,
    )
    _RUNNING_JOB_THREADS[job_id] = thread
    thread.start()
    return True


__all__ = [
    "PDF_JOB_CANCELLED_ERROR",
    "WORKER_LOST_ERROR",
    "MAX_ATTEMPTS_EXCEEDED_ERROR",
    "OWNERSHIP_LOST_ERROR",
    "WorkerOwnershipLost",
    "_RUNNING_JOB_THREADS",
    "create_translation_job",
    "create_external_excel_translation_job",
    "update_translation_job",
    "update_translation_job_owned",
    "is_translation_job_cancelled",
    "raise_if_translation_job_cancelled",
    "cancel_pdf_translation_jobs",
    "list_translation_jobs",
    "get_translation_job",
    "get_translation_job_by_id",
    "register_translation_worker",
    "heartbeat_translation_worker",
    "stop_translation_worker",
    "list_live_workers",
    "get_worker_health",
    "get_translation_queue_health",
    "get_worker_queue_health",
    "delete_translation_job",
    "list_term_candidates_for_job",
    "list_candidate_occurrences",
    "get_next_queued_pdf_job",
    "recover_stale_pdf_jobs",
    "claim_next_pdf_job",
    "claim_next_external_translation_job",
    "claim_pdf_job_by_id",
    "heartbeat_pdf_job",
    "_start_worker_heartbeat",
    "_pdf_scope_report_rows",
    "_build_excel_report_bytes",
    "run_claimed_pdf_job",
    "run_claimed_excel_job",
    "run_claimed_translation_job",
    "start_pdf_translation_job",
    "cancel_translation_jobs",
    "cancel_excel_translation_jobs",
]
